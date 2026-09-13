"""Frozen Max input snapshots and bounded-memory XLSX generation."""
from __future__ import annotations

import gzip
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

MAX_TRIALS = 1000
MAX_LINE = 8 * 1024 * 1024
MAX_UNCOMPRESSED = 512 * 1024 * 1024


def file_sha(path: Path) -> str:
    with path.open("rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest()


def write_snapshot(path: Path, *, report_run_id: str, profiles, rows, definitions, analysis_plan, segments, approved_plan, report_evidence=None) -> dict:
    if not rows or len(rows) > MAX_TRIALS:
        raise ValueError("Invalid dataset size")
    by_id = {profile.eu_number: profile for profile in profiles}
    ids = [row["trial_id"] for row in rows]
    if len(set(ids)) != len(ids) or set(ids) != set(by_id):
        raise ValueError("Dataset/profile identity mismatch")
    metadata = {"version": 1, "tier": "max", "reportRunId": report_run_id,
                "capturedAt": datetime.now(timezone.utc).isoformat(), "trialCount": len(rows),
                "definitions": definitions, "analysisPlan": analysis_plan,
                "segments": segments, "approvedPlan": approved_plan, "reportEvidence": report_evidence or []}
    with gzip.open(path, "wt", encoding="utf-8", newline="\n", compresslevel=6) as target:
        target.write(json.dumps(metadata, ensure_ascii=False, allow_nan=False) + "\n")
        for row in rows:
            if set(row["values"]) != set(definitions):
                raise ValueError("Dataset variables do not match definitions")
            record = {"row": row, "profile": by_id[row["trial_id"]].model_dump(mode="json")}
            target.write(json.dumps(record, ensure_ascii=False, allow_nan=False) + "\n")
    return {"version": 1, "status": "available", "snapshotSha": file_sha(path),
            "trialCount": len(rows), "capturedAt": metadata["capturedAt"]}


def snapshot_records(path: Path) -> Iterable[dict]:
    total = 0
    with gzip.open(path, "rt", encoding="utf-8") as source:
        while line := source.readline(MAX_LINE + 1):
            total += len(line.encode("utf-8"))
            if len(line) > MAX_LINE or total > MAX_UNCOMPRESSED:
                raise ValueError("Dataset exceeds export bounds")
            yield json.loads(line)


def fields(value: Any, pointer: str = ""):
    """JSON Pointer paths retain arrays, nulls and empty containers without loss."""
    if isinstance(value, dict):
        yield pointer, "object", None
        for key, item in value.items():
            escaped = str(key).replace("~", "~0").replace("/", "~1")
            yield from fields(item, pointer + "/" + escaped)
    elif isinstance(value, list):
        yield pointer, "array", len(value)
        for index, item in enumerate(value):
            yield from fields(item, pointer + "/" + str(index))
    elif value is None:
        yield pointer, "null", None
    elif isinstance(value, bool):
        yield pointer, "boolean", value
    elif isinstance(value, (int, float)):
        yield pointer, "number", value
    else:
        yield pointer, "string", value


def text_parts(value: str):
    # Excel limits cell text by UTF-16 units. Continuation rows preserve all text.
    part, units = [], 0
    for char in value:
        size = 2 if ord(char) > 0xFFFF else 1
        if units + size > 30000:
            yield "".join(part)
            part, units = [], 0
        part.append(char)
        units += size
    yield "".join(part)


def build_workbook(snapshot: Path, target: Path, *, expected_run_id: str) -> dict:
    # Imported only on export. Normal report generation never loads Excel tooling.
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    records = iter(snapshot_records(snapshot))
    metadata = next(records)
    if metadata.get("version") != 1 or metadata.get("tier") != "max" or metadata.get("reportRunId") != expected_run_id:
        raise ValueError("Invalid dataset identity")
    names = list(metadata["definitions"])
    if len(names) > 1000:
        raise ValueError("Too many dataset variables")
    # Validate the complete stream before opening worksheet writers. A corrupt
    # snapshot must not leave partially open XML streams or return partial data.
    validated_ids = set()
    for record in snapshot_records(snapshot):
        if "row" not in record:
            continue
        row, profile = record["row"], record["profile"]
        trial_id = row["trial_id"]
        if trial_id in validated_ids or profile["eu_number"] != trial_id or set(row["values"]) != set(names):
            raise ValueError("Dataset alignment mismatch")
        validated_ids.add(trial_id)
        if len(validated_ids) > MAX_TRIALS:
            raise ValueError("Too many dataset trials")
    if len(validated_ids) != metadata["trialCount"]:
        raise ValueError("Dataset trial count mismatch")
    workbook = Workbook(write_only=True)
    workbook.properties.title = "Max report dataset"
    workbook.properties.creator = "Trial Agents"
    header_fill = PatternFill("solid", fgColor="174B35")
    body_alignment = Alignment(vertical="top", wrap_text=True)

    class Table:
        def __init__(self, name, headings, widths):
            self.name, self.headings, self.widths = name, headings, widths
            self.part = 0
            self.open()

        def open(self):
            self.part += 1
            self.sheet = workbook.create_sheet(self.name if self.part == 1 else f"{self.name} {self.part}")
            self.sheet.freeze_panes = "B2"
            self.sheet.sheet_view.showGridLines = False
            for index, width in enumerate(self.widths, 1):
                self.sheet.column_dimensions[get_column_letter(index)].width = width
            cells = []
            for value in self.headings:
                cell = WriteOnlyCell(self.sheet, value=value)
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.fill = header_fill
                cells.append(cell)
            self.sheet.append(cells)
            self.count = 1

        def append(self, values):
            if self.count >= 1048576:
                self.finish()
                self.open()
            cells = []
            for value in values:
                if isinstance(value, int) and not isinstance(value, bool) and abs(value) > 999999999999999:
                    value = str(value)  # Preserve integers exceeding Excel's numeric precision.
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError("Nonfinite dataset number")
                if isinstance(value, float) and float(format(value, ".15g")) != value:
                    value = repr(value)
                if isinstance(value, str):
                    if ILLEGAL_CHARACTERS_RE.search(value):
                        raise ValueError("Text must be encoded before worksheet insertion")
                    if len(value.encode("utf-16-le")) > 60000:
                        raise ValueError("Long text must use continuation rows")
                cell = WriteOnlyCell(self.sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Source text is never an executable formula.
                cell.alignment = body_alignment
                cells.append(cell)
            self.sheet.append(cells)
            self.count += 1

        def finish(self):
            self.sheet.auto_filter.ref = f"A1:{get_column_letter(len(self.headings))}{self.count}"

    trials = Table("Trials", ["Trial ID", "Title", "Profile version", "Approved at", "Groups", "Uncertain groups", "Profile status"], [25, 70, 18, 28, 40, 40, 20])
    analysis = Table("Analysis data", ["Trial ID", *names], [25, *([28] * len(names))])
    variables = Table("Variables", ["Name", "Label", "Type", "Source", "Definition", "Profile field", "Analysis indices"], [28, 35, 18, 25, 75, 60, 24])
    values_table = Table("Analysis values", ["Trial ID", "Field path", "Type", "Value", "Part"], [25, 60, 22, 85, 10])
    profile_table = Table("Profile fields", ["Trial ID", "Field path", "Type", "Value", "Part"], [25, 80, 22, 85, 10])
    details = Table("Report details", ["Field path", "Type", "Value", "Part"], [65, 22, 90, 10])

    def full_fields(table, value, prefix=None):
        for pointer, kind, scalar in fields(value):
            if isinstance(scalar, str) and ILLEGAL_CHARACTERS_RE.search(scalar):
                scalar = json.dumps(scalar, ensure_ascii=True)
                kind = "string (JSON encoded)"
            parts = text_parts(scalar) if isinstance(scalar, str) else [scalar]
            for part, chunk in enumerate(parts, 1):
                table.append(([prefix] if prefix is not None else []) + [pointer, kind, chunk, part])

    def compact(value, reference):
        if isinstance(value, (list, dict)):
            return reference
        if isinstance(value, str) and (ILLEGAL_CHARACTERS_RE.search(value) or len(value.encode("utf-16-le")) > 60000):
            return reference
        return value

    for name, definition in metadata["definitions"].items():
        variables.append([name, definition.get("label"), definition.get("kind"), definition.get("source"),
                          definition.get("description"), definition.get("profile_path"), json.dumps(definition.get("analysis_indices", []))])
    full_fields(details, {**metadata, "exportNotes": {
        "scope": "Full profiles and collected variables for the trials analyzed in this report.",
        "paths": "Field paths use JSON Pointer. Array indices begin at zero. Part rows concatenate in order.",
        "missing": "The Type column distinguishes null, empty string, empty object and empty array. Null is not zero or false.",
        "analysis": "Analysis data is one row per trial. Composite or long values are complete in Analysis values.",
        "numbers": "Numbers exceeding Excel's 15-digit precision are stored as text to preserve their digits.",
        "encoding": "Only XML-incompatible strings are JSON encoded and explicitly labeled.",
    }})
    seen = set()
    for record in records:
        row, profile = record["row"], record["profile"]
        trial_id = row["trial_id"]
        if trial_id in seen or profile["eu_number"] != trial_id or set(row["values"]) != set(names):
            raise ValueError("Dataset alignment mismatch")
        seen.add(trial_id)
        if len(seen) > MAX_TRIALS:
            raise ValueError("Too many dataset trials")
        filtering = profile["profile"].get("filtering_variables", {})
        title = profile["profile"].get("classification_variables", {}).get("trial_title") or filtering.get("trial_title") or filtering.get("title") or ""
        trials.append([trial_id, compact(title, "See Profile fields"), profile.get("profile_schema_version"), profile.get("approved_at"),
                       ", ".join(row.get("segment_keys", [])), ", ".join(row.get("uncertain_segment_keys", [])), profile.get("approval_status")])
        analysis.append([trial_id, *[compact(row["values"][name], "See Analysis values") for name in names]])
        full_fields(values_table, row, trial_id)
        full_fields(profile_table, profile, trial_id)
    if len(seen) != metadata["trialCount"]:
        raise ValueError("Dataset trial count mismatch")
    for table in [trials, analysis, variables, values_table, profile_table, details]:
        table.finish()
    workbook.save(target)
    return {"trials": len(seen), "variables": len(names), "bytes": target.stat().st_size}


if __name__ == "__main__":
    import sys
    print(json.dumps(build_workbook(Path(sys.argv[1]), Path(sys.argv[2]), expected_run_id=sys.argv[3])))
