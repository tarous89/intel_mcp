"""Exercise the installed export runtime using synthetic data, with no API calls."""
import asyncio
import json
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from intel_mcp.max_agent_exports import optional_pdf, subprocess_export
from intel_mcp.max_agent_output import render_html


async def main():
    with tempfile.TemporaryDirectory(prefix="max-export-check-") as directory:
        root = Path(directory)
        payload = {
            "manifest": {"synthetic": True},
            "profiles": {"fixture": {"text": "Synthetic export check — Δ"}},
            "documents": {}, "work": {"facts": []}, "metrics": {}, "issues": [],
        }
        source, sheet = root / "input.json", root / "dataset.xlsx"
        source.write_text(json.dumps(payload), encoding="utf-8")
        await subprocess_export(source, sheet)
        book = load_workbook(sheet, read_only=True)
        try:
            values = [row[3] for row in book["Profiles"].iter_rows(min_row=2, values_only=True)]
            if json.dumps(payload["profiles"]["fixture"]["text"], ensure_ascii=False) not in values:
                raise RuntimeError("XLSX export did not preserve the synthetic input")
        finally:
            book.close()
        page, pdf = root / "report.html", root / "report.pdf"
        page.write_text(render_html(
            '<h1>Max deployment check</h1><p>Synthetic fixture — no clinical data.</p>'
            '<table><tr><th>Check</th><th>Result</th></tr><tr><td>Unicode</td><td>Δ</td></tr></table>',
            {}, "Max deployment check"), encoding="utf-8")
        if not await optional_pdf(page, pdf):
            raise RuntimeError("Sandboxed Chromium PDF check failed; worker will not claim jobs")
        print("MAX_EXPORT_CHECK passed: XLSX round-trip and sandboxed PDF; no API calls.", flush=True)


if __name__ == "__main__":
    asyncio.run(main())
