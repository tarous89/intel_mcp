import gzip
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from intel_mcp.profiles import FullProfileItem
from intel_mcp.report_dataset import build_workbook, fields, snapshot_records, text_parts, write_snapshot

RUN = "12345678-1234-1234-1234-123456789abc"


def fixture(path: Path):
    profile = FullProfileItem(eu_number="2024-123456-12-00", profile_schema_version="11.0.0", approved_at=None, approval_status="deterministic",
        profile={"filtering_variables": {"title": "=NOT_A_FORMULA", "empty": "", "null": None, "zero": 0, "false": False,
                 "long": "e\u0301😀" * 18000, "large_integer": 12345678901234567890, "control": "A\x00B"},
                 "classification_variables": {"sites": [{"name": "Site A", "investigators": [{"name": "Dr A"}, {"name": "Dr B"}]}]},
                 "empty_list": [], "empty_object": {}})
    rows = [{"trial_id": profile.eu_number, "segment_keys": ["primary"], "uncertain_segment_keys": [],
             "values": {"age": 42, "eligible": None, "investigators": ["Dr A", "Dr B"], "signal": False}}]
    definitions = {name: {"label": name, "kind": "number" if name == "age" else "categorical", "source": "profile", "description": "Definition"} for name in rows[0]["values"]}
    manifest = write_snapshot(path, report_run_id=RUN, profiles=[profile], rows=rows, definitions=definitions,
                              analysis_plan={}, segments=[], approved_plan={},
                              report_evidence=[{"supports": [{"trial_ids": [profile.eu_number], "variable_names": ["age"]}]}])
    return profile, rows, manifest


def test_snapshot_freezes_full_profiles_variables_and_definitions(tmp_path):
    source = tmp_path / "snapshot.gz"
    profile, rows, manifest = fixture(source)
    expected = profile.model_dump(mode="json")
    profile.profile["filtering_variables"]["title"] = "Changed after report"
    data = list(snapshot_records(source))
    assert data[1]["profile"] == expected
    assert data[1]["row"] == rows[0]
    assert data[0]["trialCount"] == manifest["trialCount"] == 1
    assert set(data[0]["definitions"]) == set(rows[0]["values"])
    assert data[0]["reportEvidence"][0]["supports"][0]["trial_ids"] == [profile.eu_number]


def test_workbook_preserves_values_long_text_nested_records_and_formula_text(tmp_path):
    source, target = tmp_path / "snapshot.gz", tmp_path / "dataset.xlsx"
    profile, rows, _ = fixture(source)
    result = build_workbook(source, target, expected_run_id=RUN)
    assert result["trials"] == 1 and result["variables"] == 4
    wb = load_workbook(target, read_only=True)
    assert wb.sheetnames == ["Trials", "Analysis data", "Variables", "Analysis values", "Profile fields", "Report details"]
    assert list(wb["Analysis data"].values)[1] == (profile.eu_number, 42, None, "See Analysis values", False)
    exported = list(wb["Profile fields"].values)[1:]
    long = "".join(row[3] for row in exported if row[1] == "/profile/filtering_variables/long")
    assert long == profile.profile["filtering_variables"]["long"]
    actual = {row[1]: (row[2], row[3]) for row in exported if row[4] == 1}
    assert actual["/profile/filtering_variables/null"] == ("null", None)
    assert actual["/profile/filtering_variables/zero"] == ("number", 0)
    assert actual["/profile/filtering_variables/false"] == ("boolean", False)
    assert actual["/profile/empty_list"] == ("array", 0)
    assert actual["/profile/empty_object"] == ("object", None)
    assert actual["/profile/filtering_variables/large_integer"][1] == "12345678901234567890"
    assert json.loads(actual["/profile/filtering_variables/control"][1]) == "A\x00B"
    assert actual["/profile/classification_variables/sites/0/investigators/1/name"][1] == "Dr B"
    assert {p for p, _, _ in fields(profile.model_dump(mode="json"))} == {row[1] or "" for row in exported}
    for sheet in wb:
        for row in sheet:
            assert not any(cell.data_type == "f" for cell in row)
    assert list(wb["Trials"].values)[1][1] == "=NOT_A_FORMULA"
    trial_rows = list(wb["Trials"].values)
    assert trial_rows[1][trial_rows[0].index("Profile status")] == "deterministic"
    details = {row[0]: row[2] for row in list(wb["Report details"].values)[1:]}
    assert details["/reportEvidence/0/supports/0/trial_ids/0"] == profile.eu_number
    wb.close()


def test_unicode_parts_respect_excel_limit_without_losing_characters():
    value = "😀" * 40000
    parts = list(text_parts(value))
    assert "".join(parts) == value
    assert all(len(part.encode("utf-16-le")) <= 60000 for part in parts)


def test_wrong_run_or_corrupt_cohort_is_rejected(tmp_path):
    source = tmp_path / "snapshot.gz"
    fixture(source)
    with pytest.raises(ValueError, match="identity"):
        build_workbook(source, tmp_path / "wrong.xlsx", expected_run_id="another-run")
    contents = list(snapshot_records(source))
    contents[1]["profile"]["eu_number"] = "2024-000000-00-00"
    with gzip.open(source, "wt", encoding="utf8") as out:
        for item in contents:
            out.write(json.dumps(item) + "\n")
    with pytest.raises(ValueError, match="alignment"):
        build_workbook(source, tmp_path / "invalid.xlsx", expected_run_id=RUN)
