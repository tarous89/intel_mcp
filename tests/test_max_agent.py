import asyncio
import copy
import json
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import httpx
import pytest
from lxml import html
from openpyxl import load_workbook
from starlette.applications import Starlette
from starlette.routing import Route
from starlette.testclient import TestClient

from intel_mcp.max_agent_evidence import digest, freeze_document, freeze_profile_batch, load_profiles, new_manifest, select_candidates
from intel_mcp.max_agent_exports import workbook
from intel_mcp.max_agent_output import calculate, render_html, review
from intel_mcp.max_agent_session import ManagedAgent, ManagedAgentError, scoped_token
from intel_mcp.max_agent_tools import ReportRetrieval, tool_definitions

RUN = "12345678-1234-1234-1234-123456789abc"
OTHER = "22345678-1234-1234-1234-123456789abc"
TID = "2024-000001-00-00"
PLAN = {"version": 4, "studyCohorts": [
    {"discoveryFilter": {"field": "diseases", "values": ["CRPC"]}, "selectionSegments": ["CRPC"]}
    for _ in range(3)], "reportSections": [{"sharedAnalysis": {}, "maxAnalysis": {}}]}
PROFILE = {"eu_number": TID, "profile_schema_version": "v2", "approval_status": "draft",
    "profile": {"value": 0, "disease": "CRPC", "filtering_variables": {"available_extracted_documents": {"protocol": ["Protocol.pdf"]}}}}


class MemoryEvidence:
    def __init__(self): self.values = {}
    async def put(self, run_id, value):
        sha = digest(value)
        self.values[run_id, sha] = copy.deepcopy(value)
        return sha
    async def get(self, run_id, sha): return copy.deepcopy(self.values[run_id, sha])


def test_frozen_profiles_survive_retry_and_verify_hash():
    async def scenario():
        store = MemoryEvidence()
        manifest = new_manifest(RUN, PLAN, {}, {"selectedIds": [TID]})
        first = await freeze_profile_batch(store, manifest, [PROFILE])
        retry = await freeze_profile_batch(store, first, [{**PROFILE, "approval_status": "approved"}])
        assert first["profiles"] == retry["profiles"]
        assert (await load_profiles(store, retry))[TID] == PROFILE
        sha = first["profiles"][TID]["batchSha"]
        store.values[RUN, sha]["profiles"][0]["profile"]["value"] = 999
        with pytest.raises(ValueError, match="checksum"): await load_profiles(store, first)
    asyncio.run(scenario())


def test_document_parts_are_frozen_and_cross_report_denied():
    async def scenario():
        store = MemoryEvidence()
        manifest = await freeze_profile_batch(store, new_manifest(RUN, PLAN, {}, {}), [PROFILE])
        doc = {"trial_id": TID, "document_name": "Protocol.pdf", "part": 1, "text": "[Page 1] zero", "pages": [1]}
        one = await freeze_document(store, manifest, doc)
        assert (await freeze_document(store, one, {**doc, "text": "changed"})) == one
        two = await freeze_document(store, one, {**doc, "part": 2})
        assert len(two["documents"]) == 2
        with pytest.raises(PermissionError):
            await freeze_document(store, manifest, {**doc, "trial_id": "2024-000099-00-00"})
    asyncio.run(scenario())


def test_selection_is_deterministic_and_preserves_seed_diversity():
    a = [str(n) for n in range(200)]
    b = ["subgroup-" + str(n) for n in range(20)]
    result = select_candidates([a, b, a[:20]], 100, "plan")
    assert result == select_candidates([list(reversed(a)), b, a[:20]], 100, "plan")
    assert len(result) == len(set(result)) == 100
    assert set(b) <= set(result)


def test_deduplication_denominators_zero_and_complete_rankings():
    facts = {"zero": {"trialId": TID, "value": 0}}
    members = {"group_1": {TID: True}}
    base = {"id": "m", "title": "Measured zero", "groupId": "group_1", "method": "mean",
        "observations": [{"factId": "zero", "label": "Zero"}] * 2}
    result = calculate(base, facts, members)
    assert result["rows"][0]["value"] == 0 and result["rows"][0]["n"] == 1
    assert calculate({**base, "observations": []}, facts, members)["rows"] == []
    with pytest.raises(ValueError):
        calculate({**base, "method": "percentage", "denominatorTrialIds": []}, facts, members)
    with pytest.raises(ValueError):
        calculate(base, facts, {"group_1": {TID: None}})
    many = calculate({**base, "observations": [{"factId": "zero", "label": f"Label {n}"} for n in range(15)]}, facts, members)
    doc = html.fromstring(render_html('<div data-chart="m" data-role="main"></div>', {"m": many}, "Title"))
    assert len(many["rows"]) == 15
    assert len(doc.xpath('//*[@class="chart-row"]')) == 10


def test_unsupported_sources_and_missing_coverage_are_identified():
    work = {"facts": [{"id": "bad", "trialId": TID, "value": 1, "source": {"profilePath": "/value"}}],
            "membership": {"group_1": {TID: True}}}
    _, issues = review(work, {TID: PROFILE}, {}, PLAN)
    codes = [i["code"] for i in issues]
    assert "unsupported_fact" in codes and "invalid_membership" in codes
    assert codes.count("missing_group_assessment") == 3


def test_static_html_removes_execution_and_wrong_group_main_chart():
    raw = '<script>alert(1)</script><svg><script>x</script></svg><img src="https://evil.test"><iframe srcdoc="x"></iframe><p onclick="alert(1)" style="background:url(https://evil.test)">Safe</p><a href="javascript:x">link</a><div data-chart="sub" data-role="main"></div>'
    rendered = render_html(raw, {"sub": {"groupId": "group_2"}}, "Title")
    doc = html.fromstring(rendered)
    assert not doc.xpath("//script|//svg|//iframe|//img|//*[@onclick]")
    assert "evil.test" not in rendered and "javascript:x" not in rendered
    assert "Safe" in rendered and "default-src" in rendered


def test_workbook_full_rankings_lossless_text_and_no_formulas(tmp_path):
    text = "=HYPERLINK(\"evil\")" + "😀" * 20000 + "\u0001"
    payload = {"manifest": {"version": 2}, "profiles": {TID: {"text": text}},
        "documents": {}, "work": {"facts": []}, "metrics": {"rows": list(range(15))}, "issues": [], "scripts": {}}
    first, second = tmp_path / "one.xlsx", tmp_path / "two.xlsx"
    workbook(payload, first); workbook(payload, second)
    assert first.read_bytes() == second.read_bytes()
    book = load_workbook(first, read_only=True)
    rows = list(book["Profiles"].iter_rows(min_row=2))
    cells = [r[3] for r in rows if r[0].value == "/" + TID + "/text"]
    assert len(cells) > 1 and all(c.data_type == "s" for c in cells)
    assert json.loads("".join(c.value for c in cells)) == text
    assert len(list(book["Rankings and results"].rows)) == 18
    book.close()


def test_scoped_mcp_listing_dispatch_and_report_auth(monkeypatch):
    import intel_mcp.max_agent_tools as module
    calls = []
    async def control(action, **payload):
        if payload["reportRunId"] != RUN: raise PermissionError()
        return {}
    retrieval = NS(control=NS(call=control), call=AsyncMock(return_value={"ok": True}))
    monkeypatch.setattr(module, "ReportRetrieval", lambda _: retrieval)
    class MCP:
        routes = []
        def custom_route(self, path, methods):
            def register(fn):
                self.routes.append(Route(path, fn, methods=methods))
                return fn
            return register
    mcp = MCP()
    settings = NS(app_service_token="test-signing-key")
    module.register_max_tools(mcp, settings)
    with TestClient(Starlette(routes=mcp.routes)) as client:
        headers = {"Authorization": "Bearer " + scoped_token(settings.app_service_token, RUN)}
        def call(name, params=None):
            return client.post(f"/max-agent/mcp/{RUN}", json={"jsonrpc": "2.0", "id": 1, "method": name, "params": params or {}}, headers=headers)
        listing = call("tools/list").json()
        assert [t["name"] for t in listing["result"]["tools"]] == ["filter_trials", "get_profiles", "get_documents"]
        for forbidden in ["extract_variables", "classify_trials", "start_analysis", "admin", "worker"]:
            assert "error" in call("tools/call", {"name": forbidden}).json()
        assert retrieval.call.await_count == 0
        assert client.post(f"/max-agent/mcp/{OTHER}", json={"method": "tools/list"}, headers=headers).status_code == 401
        assert client.post(f"/max-agent/mcp/{RUN}", json={"method": "tools/list"}).status_code == 401


def test_snapshot_reads_do_not_hit_live_engine():
    async def scenario():
        store = MemoryEvidence()
        manifest = await freeze_profile_batch(store, new_manifest(RUN, PLAN, {}, {}), [PROFILE])
        sha = await store.put(RUN, manifest)
        control = NS(call=AsyncMock(return_value={"job": {"id": RUN, "base_version": 1,
            "request_text": "Change formatting", "checkpoint": {"evidenceSha": sha}}, "analysisId": "lease"}))
        engine = NS(get_profiles=AsyncMock(), filter_trials=AsyncMock())
        r = ReportRetrieval(NS(), control=control, evidence=store, engine=engine, allowances=NS(), config=NS())
        assert (await r.call(RUN, "get_profiles", {"trial_ids": [TID]}))["profiles"] == [PROFILE]
        with pytest.raises(PermissionError):
            await r.call(RUN, "filter_trials", {"live": True})
        engine.get_profiles.assert_not_awaited(); engine.filter_trials.assert_not_awaited()
    asyncio.run(scenario())


def test_managed_session_is_idle_single_agent_and_events_idempotent():
    async def scenario():
        sent = []
        def handler(request):
            sent.append(request)
            return httpx.Response(200, json={"id": "session-1"})
        settings = NS(openai_api_key="test", openai_base_url="https://api.openai.com/v1",
                      mcp_public_resource_url="https://mcp.example/mcp", app_service_token="scope")
        config = NS(model="pilot-model", validate=lambda _: None)
        agent = ManagedAgent(settings, config, transport=httpx.MockTransport(handler))
        await agent.create_idle(RUN, "intent", "instructions", {"evidence.json.gz": b"data"})
        body = json.loads(sent[0].content)
        assert "input" not in body
        assert body["agent"]["multi_agent"] == {"enabled": False}
        assert body["agent"]["tools"][0]["allowed_tools"] == ["filter_trials", "get_profiles", "get_documents"]
        assert body["environment"]["network"]["access"] == "disabled"
        assert b"test" not in json.dumps(body["environment"]).encode()
        await agent.send("session-1", "request", "stable")
        await agent.send("session-1", "request", "stable")
        assert sent[1].headers["idempotency-key"] == sent[2].headers["idempotency-key"] == "stable"
    asyncio.run(scenario())


def test_evidence_worker_resumes_without_refetching_completed_batch():
    from intel_mcp.max_agent_execution import MaxAgentWorker
    async def scenario():
        store = MemoryEvidence()
        manifest = await freeze_profile_batch(store, new_manifest(RUN, PLAN, {}, {"selectedIds": [TID]}), [PROFILE])
        sha = await store.put(RUN, manifest)
        worker = MaxAgentWorker(NS(), control=NS(call=AsyncMock()), engine=NS(get_profiles=AsyncMock()),
            evidence=store, agent=NS(), allowances=NS())
        job = {"id": RUN, "report_run_id": RUN, "claim_token": OTHER, "checkpoint": {"evidenceSha": sha}, "status": "preparing"}
        await worker.prepare(job, {"plan": PLAN}, {"analysisId": "lease"})
        assert job["status"] == "analyzing"
        worker.engine.get_profiles.assert_not_awaited()
    asyncio.run(scenario())


def test_malformed_optional_data_retains_valid_results_without_mutating_work():
    work = {
        "facts": [{"id": "zero", "trialId": TID, "value": 0, "source": {"profilePath": "/value"}}],
        "membership": {"group_2": {TID: True}, "group_1": {TID: True}, "group_3": [TID]},
        "membershipSources": {"group_1": {TID: ["zero"]}, "group_2": {TID: ["zero"]}},
        "metrics": [{"id": "valid", "analysisId": "analysis_1", "groupId": "group_1",
            "method": "mean", "observations": [{"factId": "zero", "label": "Zero"}]}, None],
        "assessments": [None],
    }
    original = copy.deepcopy(work)
    metrics, issues = review(work, {TID: PROFILE}, {}, PLAN)
    assert metrics["valid"]["rows"][0]["value"] == 0
    assert work == original
    assert {i["code"] for i in issues} >= {"invalid_membership", "unsupported_metric", "missing_group_assessment"}
    for bad in [None, [], {"facts": None, "membership": None, "metrics": None, "assessments": None}]:
        _, issues = review(bad, {TID: PROFILE}, {}, PLAN)
        assert issues


def test_large_input_uses_file_reference_without_truncation():
    async def scenario():
        sent = []
        def handler(request):
            sent.append(request)
            return httpx.Response(200, json={"id": "file-evidence" if request.url.path.endswith("/files") else "session"})
        settings = NS(openai_api_key="test", openai_base_url="https://api.openai.com/v1",
                      mcp_public_resource_url="https://mcp.example/mcp", app_service_token="scope")
        agent = ManagedAgent(settings, NS(model="pilot", validate=lambda _: None), transport=httpx.MockTransport(handler))
        data = b"x" * (6 * 1024 * 1024)
        file_id = await agent.upload_input("evidence.json.gz", data)
        assert data in sent[0].content
        await agent.create_idle(RUN, "intent", "instructions", {"evidence.json.gz": data}, {"evidence.json.gz": file_id})
        body = json.loads(sent[1].content)
        assert body["environment"]["files"] == [{"type": "file_id", "path": "/workspace/evidence.json.gz", "file_id": "file-evidence"}]
    asyncio.run(scenario())


def test_restored_session_does_not_publish_an_older_turn():
    from intel_mcp.max_agent_execution import MaxAgentWorker
    async def scenario():
        agent = NS(inspect=AsyncMock(return_value=({"status": "idle"}, [
            {"id": "last", "status": "completed"}, {"id": "older", "status": "completed"}])), send=AsyncMock())
        worker = MaxAgentWorker(NS(), control=NS(call=AsyncMock()), engine=NS(),
            evidence=MemoryEvidence(), agent=agent, allowances=NS())
        job = {"id": RUN, "report_run_id": RUN, "claim_token": OTHER, "base_version": 1,
            "session_id": "saved-session", "checkpoint": {"turnIntent": "current-request", "previousTurnId": "last"}, "status": "analyzing"}
        await worker.analyze(job)
        agent.send.assert_awaited_once()
        assert job["status"] == "analyzing"
        assert not job.get("turn_id")
        await worker.checkpoint(job, "analyzing", sessionId="")
        assert job["session_id"] == ""
    asyncio.run(scenario())


def test_pdf_retry_uses_published_html_without_agent_work(monkeypatch):
    from intel_mcp import max_agent_pdf as module
    async def scenario():
        calls = []
        job = {"report_run_id": RUN, "version": 2, "claim_token": OTHER, "html_sha": "a" * 64}
        async def call(action, **payload):
            calls.append((action, payload))
            return {"job": job} if action == "claim_pdf" else {"ok": True}
        async def download(run, sha, extension, path):
            assert (run, sha, extension) == (RUN, "a" * 64, "html")
            path.write_text("<h1>Published version 2</h1>")
        async def render(source, target):
            assert source.read_text() == "<h1>Published version 2</h1>"
            target.write_bytes(b"%PDF-test")
            return True
        monkeypatch.setattr(module, "optional_pdf", render)
        artifacts = NS(download=download, upload=AsyncMock(return_value="b" * 64))
        assert await module.retry_pdf(NS(call=call), artifacts)
        assert calls[-1] == ("complete_pdf", {"reportRunId": RUN, "version": 2, "claimToken": OTHER, "pdfSha": "b" * 64})
        monkeypatch.setattr(module, "optional_pdf", AsyncMock(return_value=False))
        assert await module.retry_pdf(NS(call=call), artifacts)
        assert calls[-1][0] == "retry_pdf"
    asyncio.run(scenario())
